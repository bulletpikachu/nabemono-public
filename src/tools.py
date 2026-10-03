import asyncio
import io
import json
import re
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlparse

import aiohttp
import discord

from src.config import (
    BRAVE_API_KEY,
    BRAVE_MAX_RESULTS,
    BRAVE_TIMEOUT,
    CHANNEL_HISTORY_DEFAULT_LIMIT,
    CHANNEL_HISTORY_MAX_CHARS,
    CHANNEL_HISTORY_MAX_RESULTS,
    CHANNEL_HISTORY_MAX_SCAN,
    SEND_FILE_MAX_CHARS,
    VISIT_PAGE_MAX_BYTES,
    VISIT_PAGE_MAX_CHARS,
    VISIT_PAGE_TIMEOUT,
)

SNIPPET_LIMIT = 300
BRAVE_SEARCH_URL = "https://api.search.brave.com/res/v1/web/search"
BRAVE_FRESHNESS = {
    "day": "pd",
    "week": "pw",
    "month": "pm",
    "year": "py",
}


async def reminder(original_message, reminder_message, **schedule):
    print("[TOOL] reminder")
    from src.reminders import create_reminder

    schedule.pop("dm", None)
    return json.dumps(
        create_reminder(original_message, reminder_message, **schedule),
        ensure_ascii=False,
    )


async def list_reminders(original_message):
    print("[TOOL] list_reminders")
    from src.reminders import list_reminders as list_user_reminders

    return json.dumps(list_user_reminders(original_message), ensure_ascii=False)


async def cancel_reminder(original_message, reminder_id):
    print("[TOOL] cancel_reminder")
    from src.reminders import cancel_reminder as cancel_user_reminder

    return json.dumps(
        cancel_user_reminder(original_message, reminder_id),
        ensure_ascii=False,
    )


async def canvas_upcoming(original_message, days=None):
    print("[TOOL] canvas_upcoming")
    from src.canvas.service import canvas_service, upcoming_payload

    result = await canvas_service.upcoming(
        original_message.author.id,
        days=days,
    )
    return json.dumps(upcoming_payload(result), ensure_ascii=False)


async def interim_message(original_message, message):
    print("[TOOL] interim_message")
    text = (message or "").strip()
    if not text:
        raise ValueError("message must not be empty")
    if len(text) > 1900:
        text = text[:1900]

    sent = await original_message.channel.send(text)
    created_at = sent.created_at or datetime.now(timezone.utc)
    if created_at.tzinfo is None:
        created_at = created_at.replace(tzinfo=timezone.utc)

    return json.dumps(
        {
            "status": "sent",
            "message_id": sent.id,
            "channel_id": original_message.channel.id,
            "created_at": created_at.isoformat(),
        },
        ensure_ascii=False,
    )


def _clamp_int(value, default, minimum, maximum):
    if value is None:
        value = default
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValueError("limit must be an integer")
    return max(minimum, min(maximum, number))


def _history_anchor(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        if value.tzinfo is None:
            return value.replace(tzinfo=timezone.utc)
        return value
    if isinstance(value, int):
        return discord.Object(id=value)
    text = str(value).strip()
    if not text:
        return None
    if text.isdigit() and len(text) >= 17:
        return discord.Object(id=int(text))
    parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed


def _author_name(author):
    return (
        getattr(author, "display_name", None)
        or getattr(author, "global_name", None)
        or getattr(author, "name", None)
        or str(getattr(author, "id", "unknown"))
    )


def _author_matches(message, author):
    if not author:
        return True
    needle = str(author).strip().lower()
    if not needle:
        return True
    user = message.author
    candidates = [
        str(getattr(user, "id", "")),
        getattr(user, "name", "") or "",
        getattr(user, "display_name", "") or "",
        getattr(user, "global_name", "") or "",
    ]
    return any(candidate.lower() == needle for candidate in candidates if candidate)


def _message_matches_query(message, query):
    if not query:
        return True
    needle = str(query).strip().lower()
    if not needle:
        return True
    content = (message.content or "").lower()
    author = _author_name(message.author).lower()
    return needle in content or needle in author


def format_channel_history(
    channel_id,
    messages,
    truncated=False,
    scanned=0,
    query=None,
    max_chars=CHANNEL_HISTORY_MAX_CHARS,
):
    remaining = max_chars
    formatted = []
    content_truncated = False
    for message in messages:
        created_at = message.created_at or datetime.now(timezone.utc)
        if created_at.tzinfo is None:
            created_at = created_at.replace(tzinfo=timezone.utc)
        content = message.content or ""
        entry = {
            "id": message.id,
            "author": _author_name(message.author),
            "created_at": created_at.isoformat(),
            "content": content,
        }
        reference = getattr(message, "reference", None)
        reply_id = getattr(reference, "message_id", None) if reference else None
        if reply_id:
            entry["reply_to"] = reply_id
        attachments = [
            getattr(attachment, "filename", None) or "attachment"
            for attachment in (getattr(message, "attachments", None) or [])
        ]
        if attachments:
            entry["attachments"] = attachments

        encoded = json.dumps(entry, ensure_ascii=False)
        if formatted and len(encoded) > remaining:
            content_truncated = True
            break
        if len(encoded) > remaining:
            overflow = len(encoded) - remaining
            if overflow < len(content):
                entry["content"] = content[: len(content) - overflow].rstrip() + "..."
                encoded = json.dumps(entry, ensure_ascii=False)
            content_truncated = True
            if len(encoded) > remaining:
                break
        formatted.append(entry)
        remaining -= len(encoded)

    summary = {
        "channel_id": channel_id,
        "count": len(formatted),
        "scanned": scanned,
        "messages": formatted,
    }
    if query:
        summary["query"] = query
    if truncated or content_truncated:
        summary["truncated"] = True
    if not formatted:
        summary["note"] = "no matching messages found"
    return summary


async def read_channel_history(
    original_message,
    limit=None,
    query=None,
    before=None,
    after=None,
    author=None,
):
    print("[TOOL] read_channel_history")
    channel = getattr(original_message, "channel", None)
    if channel is None or not hasattr(channel, "history"):
        raise ValueError("read_channel_history can only be used in a Discord channel")

    result_limit = _clamp_int(
        limit,
        CHANNEL_HISTORY_DEFAULT_LIMIT,
        1,
        CHANNEL_HISTORY_MAX_RESULTS,
    )
    query_text = (query or "").strip() or None
    author_text = (author or "").strip() or None
    before_anchor = _history_anchor(before)
    after_anchor = _history_anchor(after)
    scanning = bool(query_text or author_text)
    fetch_limit = CHANNEL_HISTORY_MAX_SCAN if scanning else result_limit

    matched = []
    scanned = 0
    try:
        history = channel.history(
            limit=fetch_limit,
            before=before_anchor,
            after=after_anchor,
        )
        async for message in history:
            scanned += 1
            if not _author_matches(message, author_text):
                continue
            if not _message_matches_query(message, query_text):
                continue
            matched.append(message)
            if len(matched) >= result_limit:
                break
    except Exception as error:
        print(f"[TOOL] read_channel_history error: {error}")
        return json.dumps(
            {
                "channel_id": getattr(channel, "id", None),
                "error": str(error),
                "note": "could not read this channel's history.",
            },
            ensure_ascii=False,
        )

    matched.reverse()
    truncated = len(matched) >= result_limit or (
        scanning and scanned >= fetch_limit
    )
    return json.dumps(
        format_channel_history(
            getattr(channel, "id", None),
            matched,
            truncated=truncated,
            scanned=scanned,
            query=query_text,
        ),
        ensure_ascii=False,
    )


SEND_FILE_FORMATS = {"txt", "docx"}


def _sanitize_filename(filename, fmt):
    name = Path(str(filename or "file")).name.strip()
    name = re.sub(r"[^\w.\- ]+", "_", name).strip(" ._")
    stem = Path(name).stem.strip() or "file"
    stem = stem[:80]
    return f"{stem}.{fmt}"


def build_text_file(content):
    return content.encode("utf-8")


def build_docx_file(content):
    from docx import Document

    document = Document()
    text = content or ""
    if text:
        for line in text.split("\n"):
            document.add_paragraph(line)
    else:
        document.add_paragraph("")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def send_file(original_message, filename, content, format="txt", caption=""):
    print("[TOOL] send_file")
    channel = getattr(original_message, "channel", None)
    if channel is None or not hasattr(channel, "send"):
        raise ValueError("send_file can only be used in a Discord channel")

    fmt = str(format or "txt").strip().lower().lstrip(".")
    if fmt not in SEND_FILE_FORMATS:
        raise ValueError("format must be txt or docx")
    if not isinstance(content, str) or not content.strip():
        raise ValueError("content must be a non-empty string")
    if len(content) > SEND_FILE_MAX_CHARS:
        raise ValueError(
            f"content exceeds the {SEND_FILE_MAX_CHARS} character limit"
        )

    safe_name = _sanitize_filename(filename, fmt)
    if fmt == "docx":
        payload = build_docx_file(content)
    else:
        payload = build_text_file(content)

    caption_text = (caption or "").strip()
    if len(caption_text) > 1900:
        caption_text = caption_text[:1900]

    send_kwargs = {
        "file": discord.File(io.BytesIO(payload), filename=safe_name),
    }
    if caption_text:
        send_kwargs["content"] = caption_text

    sent = await channel.send(**send_kwargs)
    created_at = sent.created_at or datetime.now(timezone.utc)
    if created_at.tzinfo is None:
        created_at = created_at.replace(tzinfo=timezone.utc)

    return json.dumps(
        {
            "status": "sent",
            "filename": safe_name,
            "format": fmt,
            "bytes": len(payload),
            "message_id": sent.id,
            "channel_id": getattr(channel, "id", None),
            "created_at": created_at.isoformat(),
        },
        ensure_ascii=False,
    )


def _trim_snippet(content):
    content = content or ""
    if len(content) <= SNIPPET_LIMIT:
        return content
    return content[:SNIPPET_LIMIT].rstrip() + "..."


def format_search_results(payload, max_results=BRAVE_MAX_RESULTS):
    web_results = (payload.get("web") or {}).get("results") or []
    query = payload.get("query") or {}
    summary = {
        "query": query.get("original", ""),
        "results": [
            {
                "title": result.get("title", ""),
                "url": result.get("url", ""),
                "content": _trim_snippet(result.get("description", "")),
            }
            for result in web_results[:max_results]
        ],
    }

    if not summary["results"]:
        summary["note"] = "no results returned"

    return summary


async def web_search(original_message, query, time_range=None):
    print("[TOOL] web_search")
    if not BRAVE_API_KEY:
        raise RuntimeError("BRAVE_API_KEY is not configured")

    params = {"q": query, "count": BRAVE_MAX_RESULTS}
    if time_range:
        params["freshness"] = BRAVE_FRESHNESS.get(time_range, time_range)

    headers = {
        "Accept": "application/json",
        "Accept-Encoding": "gzip",
        "X-Subscription-Token": BRAVE_API_KEY,
    }
    timeout = aiohttp.ClientTimeout(total=BRAVE_TIMEOUT)
    async with aiohttp.ClientSession(timeout=timeout) as session:
        async with session.get(
            BRAVE_SEARCH_URL,
            headers=headers,
            params=params,
        ) as response:
            response.raise_for_status()
            payload = await response.json(content_type=None)

    return json.dumps(format_search_results(payload), ensure_ascii=False)


class _HTMLTextExtractor(HTMLParser):
    SKIP_TAGS = {
        "script",
        "style",
        "noscript",
        "svg",
        "template",
        "iframe",
        "nav",
        "footer",
        "header",
    }
    BLOCK_TAGS = {
        "p",
        "div",
        "br",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "li",
        "tr",
        "section",
        "article",
        "blockquote",
        "pre",
    }

    def __init__(self):
        super().__init__()
        self._skip_depth = 0
        self._in_title = False
        self.title_parts = []
        self.chunks = []

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP_TAGS:
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if tag == "title":
            self._in_title = True
        if tag in self.BLOCK_TAGS:
            self.chunks.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP_TAGS and self._skip_depth:
            self._skip_depth -= 1
            return
        if self._skip_depth:
            return
        if tag == "title":
            self._in_title = False
        if tag in self.BLOCK_TAGS:
            self.chunks.append("\n")

    def handle_data(self, data):
        if self._skip_depth:
            return
        text = data.strip()
        if not text:
            return
        if self._in_title:
            self.title_parts.append(text)
        else:
            self.chunks.append(text)


def extract_page_text(html):
    parser = _HTMLTextExtractor()
    parser.feed(html or "")
    title = " ".join(parser.title_parts).strip()
    lines = [
        re.sub(r"\s+", " ", line).strip()
        for line in "".join(parser.chunks).splitlines()
    ]
    content = "\n".join(line for line in lines if line)
    return title, content


def format_page_contents(
    url,
    title,
    content,
    content_type="",
    max_chars=VISIT_PAGE_MAX_CHARS,
):
    content = content or ""
    truncated = len(content) > max_chars
    if truncated:
        content = content[:max_chars].rstrip() + "..."

    summary = {
        "url": url,
        "title": title,
        "content": content,
    }
    if content_type:
        summary["content_type"] = content_type
    if truncated:
        summary["truncated"] = True
    if not content:
        summary["note"] = "no readable text found"
    return summary


def _normalize_url(url):
    parsed = urlparse((url or "").strip())
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise ValueError("url must be an http or https address")
    return parsed.geturl()


def _page_text_from_body(body, content_type):
    media_type = (content_type or "").split(";", 1)[0].strip().lower()
    if media_type in {"text/html", "application/xhtml+xml"} or (
        not media_type and body.lstrip().startswith("<")
    ):
        return extract_page_text(body)
    if media_type.startswith("text/") or media_type in {
        "application/json",
        "application/xml",
        "application/javascript",
    }:
        return "", body.strip()
    return "", ""


async def visit_page(original_message, url):
    print("[TOOL] visit_page")
    url = _normalize_url(url)
    headers = {
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,text/plain;q=0.8,*/*;q=0.1",
        "Accept-Language": "en-US,en;q=0.9",
        "User-Agent": "Mozilla/5.0 (compatible; nabemono/1.0)",
    }
    timeout = aiohttp.ClientTimeout(total=VISIT_PAGE_TIMEOUT)
    async with aiohttp.ClientSession(timeout=timeout) as session:
        async with session.get(url, headers=headers, allow_redirects=True) as response:
            final_url = str(response.url)
            if response.status >= 400:
                reason = (response.reason or "").strip()
                print(f"[TOOL] visit_page HTTP {response.status} {final_url}")
                return json.dumps(
                    {
                        "url": final_url,
                        "error": f"HTTP {response.status}{f' {reason}' if reason else ''}",
                        "note": "the page refused or failed the request. try a different URL.",
                    },
                    ensure_ascii=False,
                )
            content_type = response.headers.get("Content-Type", "")
            raw = await response.content.read(VISIT_PAGE_MAX_BYTES)
            charset = response.charset or "utf-8"
            body = raw.decode(charset, errors="replace")

    title, content = _page_text_from_body(body, content_type)
    return json.dumps(
        format_page_contents(final_url, title, content, content_type),
        ensure_ascii=False,
    )


async def maps_search(original_message, query, latitude=None, longitude=None):
    print("[TOOL] maps_search")
    from src.maps import search_maps

    return json.dumps(
        await search_maps(query, latitude, longitude),
        ensure_ascii=False,
    )


async def wikipedia_search(original_message, query):
    print("[TOOL] wikipedia_search")
    from src.wikipedia import search_pages

    return json.dumps(search_pages(query), ensure_ascii=False)


async def wikipedia_page(original_message, title):
    print("[TOOL] wikipedia_page")
    from src.wikipedia import get_page

    return json.dumps(get_page(title), ensure_ascii=False)


async def music_queue(original_message):
    print("[TOOL] music_queue")
    from src.music.tools import music_queue as get_music_queue

    return await get_music_queue(original_message)


async def youtube_search(original_message, query, max_results=5):
    print("[TOOL] youtube_search")
    from src.music.tools import youtube_search as search_youtube

    return await search_youtube(original_message, query, max_results)


async def music_play(original_message, url):
    print("[TOOL] music_play")
    from src.music.tools import music_play as play_music

    return await play_music(original_message, url)


async def persona_search(original_message, query):
    print("[TOOL] persona_search")
    from src.persona.service import persona_store

    matches = await asyncio.to_thread(persona_store.search, query)
    return json.dumps(
        {
            "query": str(query or "").strip(),
            "found": bool(matches),
            "matches": [
                {
                    "key": fact.key,
                    "value": fact.value,
                    "match": fact.match,
                }
                for fact in matches
            ],
            "note": (
                "These facts are authoritative."
                if matches
                else "No matching persona fact exists. Do not invent an answer."
            ),
        },
        ensure_ascii=False,
    )


# Map function names to implementations
available_functions = {
    "reminder": reminder,
    "list_reminders": list_reminders,
    "cancel_reminder": cancel_reminder,
    "canvas_upcoming": canvas_upcoming,
    "interim_message": interim_message,
    "read_channel_history": read_channel_history,
    "send_file": send_file,
    "web_search": web_search,
    "maps_search": maps_search,
    "visit_page": visit_page,
    "wikipedia_search": wikipedia_search,
    "wikipedia_page": wikipedia_page,
    "music_queue": music_queue,
    "youtube_search": youtube_search,
    "music_play": music_play,
    "persona_search": persona_search,
}

async def execute_tool_call(original_message, tool_call):
    function_to_call = available_functions[tool_call.name]
    return await function_to_call(original_message, **(tool_call.arguments or {}))
